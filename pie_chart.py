import openpyxl
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, PieChart, ScatterChart
from openpyxl.chart.label import DataLabelList

wb = Workbook()
ws = wb.active
ws.title = "Monthly Sales Chart"

df = pd.read_csv('demo_data.csv')
data = np.vstack([df.columns.values, df.values])

for row in data:
    ws.append(row.tolist())

chart_row = 1

pie_sheet = wb.create_sheet("Pie Charts")
pie_chart = PieChart()
pie_chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(data)), titles_from_data=True)
pie_chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(data)))
pie_chart.title = "Pie Chart Sales"
pie_sheet.add_chart(pie_chart, f"A{chart_row}")

pie_chart.dataLabels = DataLabelList()
# pie_chart.dataLabels.showVal = True
pie_chart.dataLabels.showPercent = True

pie_chart_1 = PieChart()
pie_chart_1.add_data(Reference(ws, min_col=3, min_row=2, max_row=len(data)))
pie_chart_1.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(data)))
pie_chart_1.title = "Pie Chart Expenses"
pie_sheet.add_chart(pie_chart_1, f"A{chart_row + 20}")

wb.save("pie_chart.xlsx")
