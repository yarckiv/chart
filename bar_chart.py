import openpyxl
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, PieChart, ScatterChart
from openpyxl.chart.label import DataLabelList


wb = Workbook()
ws = wb.active
ws.title = "Monthly Sales Chart"

wb = Workbook()
ws = wb.active
ws.title = "Monthly Sales Chart"

df = pd.read_csv('demo_data.csv')
data = np.vstack([df.columns.values, df.values])

for row in data:
    ws.append(row.tolist())

bar_chart = BarChart()
bar_chart.type = "col"
bar_chart.style = 10
bar_chart.title = "Monthly Sales Performance"
bar_chart.y_axis.title = "Sales (USD)"
bar_chart.x_axis.title = "Month"

data_ref = Reference(ws, min_col=2, min_row=1, max_col=len(data[0]), max_row=len(data))
categories_ref = Reference(ws, min_col=1, min_row=2, max_row=len(data))

bar_chart.data_labels = DataLabelList()
bar_chart.data_labels.showVal = True

bar_chart.add_data(data_ref, titles_from_data=True)
bar_chart.set_categories(categories_ref)

ws.add_chart(bar_chart, "G2")

wb.save("bar_chart.xlsx")

